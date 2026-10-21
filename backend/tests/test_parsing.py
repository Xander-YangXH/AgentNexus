"""文件解析（阶段 2.1）。"""

import io

import pytest

from agentnexus.knowledge.parsing import ParseError, Section, parse


def make_pdf(pages: list[str]) -> bytes:
    """手工拼一个最小的合法 PDF（Helvetica、每页一行英文），避免往仓库里放二进制文件。"""
    objects = ["<< /Type /Catalog /Pages 2 0 R >>", None, "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>"]
    kids = []
    for text in pages:
        content = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode()
        objects.append(f"<< /Length {len(content)} >>\nstream\n{content.decode()}\nendstream")
        content_ref = len(objects)
        objects.append(f"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents {content_ref} 0 R /Resources << /Font << /F1 3 0 R >> >> >>")
        kids.append(f"{len(objects)} 0 R")
    objects[1] = f"<< /Type /Pages /Kids [{' '.join(kids)}] /Count {len(kids)} >>"
    out, offsets = io.BytesIO(), []
    out.write(b"%PDF-1.4\n")
    for i, obj in enumerate(objects, start=1):
        offsets.append(out.tell())
        out.write(f"{i} 0 obj\n{obj}\nendobj\n".encode())
    xref = out.tell()
    out.write(f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode())
    for off in offsets:
        out.write(f"{off:010d} 00000 n \n".encode())
    out.write(f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n".encode())
    return out.getvalue()


def test_text_utf8_with_bom_and_gbk():
    assert parse("a.txt", "﻿退货政策".encode("utf-8")) == [Section("退货政策")]
    assert parse("a.md", "退货政策".encode("gbk")) == [Section("退货政策")]  # Windows 上常见的 GBK 文件


def test_pdf_pages_keep_page_numbers():
    sections = parse("manual.pdf", make_pdf(["Return policy page one", "Refund timing page two"]))
    assert [(s.page, s.text) for s in sections] == [(1, "Return policy page one"), (2, "Refund timing page two")]


def test_docx_paragraphs_and_tables():
    from docx import Document

    doc = Document()
    doc.add_paragraph("会员规则")
    doc.add_paragraph("每消费 1 元获得 1 积分。")
    table = doc.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text, table.rows[0].cells[1].text = "等级", "折扣"
    buf = io.BytesIO()
    doc.save(buf)
    [section] = parse("rules.docx", buf.getvalue())
    assert "每消费 1 元获得 1 积分。" in section.text and "等级 | 折扣" in section.text  # 表格内容也要拿到


@pytest.mark.parametrize(
    "filename, data, code",
    [
        ("a.exe", b"MZ", "unsupported_type"),
        ("noext", b"x", "unsupported_type"),
        ("a.txt", b"   \n\t ", "empty"),
        ("a.pdf", b"not a pdf", "corrupt"),
        ("a.docx", b"not a zip", "corrupt"),
        ("a.txt", b"\xff\xfe\xfa\xfb\xfc", "corrupt"),  # 既不是 UTF-8 也不是 GB18030
    ],
)
def test_errors_are_classified(filename, data, code):
    with pytest.raises(ParseError) as e:
        parse(filename, data)
    assert e.value.code == code


def test_pdf_without_text_is_empty_not_success():
    with pytest.raises(ParseError) as e:
        parse("scan.pdf", make_pdf([""]))  # 没有文字层（类似扫描件）
    assert e.value.code == "empty"
