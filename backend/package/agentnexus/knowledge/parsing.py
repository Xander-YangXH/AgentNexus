"""把上传的文件解析成纯文本（阶段 2，decisions/0014）。

支持：.txt / .md（按 UTF-8 解码，失败再按 GB18030，覆盖 Windows 上常见的 GBK 中文文件）、.pdf（pypdf，逐页）、.docx（python-docx）。
不支持的类型、加密 PDF、解析不出任何文字（如扫描件，需要 OCR）都明确报错，而不是入库一个空文档。
"""

import io
from dataclasses import dataclass
from pathlib import PurePosixPath


@dataclass(frozen=True)
class Section:
    text: str
    page: int | None = None  # PDF 的页码（从 1 开始），其他格式为 None；用于回答时标注引用出处


class ParseError(Exception):
    def __init__(self, code: str, message: str):
        super().__init__(message)
        self.code = code  # unsupported_type / encrypted / empty / corrupt / parse_timeout / too_complex / parse_failed


SUPPORTED_EXTENSIONS = {".txt", ".md", ".markdown", ".pdf", ".docx"}


def extension_of(filename: str) -> str:
    return PurePosixPath(filename.lower()).suffix


def parse(filename: str, data: bytes) -> list[Section]:
    """按扩展名选择解析器；返回非空的段落列表。"""
    ext = extension_of(filename)
    if ext not in SUPPORTED_EXTENSIONS:
        raise ParseError("unsupported_type", f"不支持的文件类型：{ext or '（无扩展名）'}")
    try:
        if ext == ".pdf":
            sections = _parse_pdf(data)
        elif ext == ".docx":
            sections = _parse_docx(data)
        else:
            sections = [Section(_decode_text(data))]
    except ParseError:
        raise
    except Exception as exc:  # 文件损坏、格式不对等：统一归为 corrupt，原始信息留在消息里供排查
        raise ParseError("corrupt", f"文件无法解析：{type(exc).__name__}: {exc}") from exc
    sections = [s for s in (Section(s.text.strip(), s.page) for s in sections) if s.text]
    if not sections:
        raise ParseError("empty", "没有解析出任何文字（如果是扫描件，需要先做 OCR）")
    return sections


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030"):  # utf-8-sig 会顺带去掉 BOM
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("corrupt", "无法识别文本编码（只支持 UTF-8 与 GBK/GB18030）")


def _parse_pdf(data: bytes) -> list[Section]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        raise ParseError("encrypted", "PDF 已加密，无法读取")
    return [Section(page.extract_text() or "", page=i) for i, page in enumerate(reader.pages, start=1)]


def _parse_docx(data: bytes) -> list[Section]:
    from docx import Document

    doc = Document(io.BytesIO(data))
    parts = [p.text for p in doc.paragraphs]
    for table in doc.tables:  # 表格里的文字也很常见（如价格表、参数表），按“单元格 | 单元格”拼成行
        for row in table.rows:
            parts.append(" | ".join(cell.text.strip() for cell in row.cells))
    return [Section("\n".join(parts))]
